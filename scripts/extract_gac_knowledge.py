"""Read-only extraction of workbook observations and guidance, with cell provenance."""
import argparse
import hashlib
import json
from pathlib import Path
import openpyxl

def numeric(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None

def observation(value):
    if numeric(value) is not None:
        return {"value": numeric(value), "status": "observed", "raw": value}
    text = str(value or "").strip()
    status = "not_monitored" if not text else "not_reached" if text == "no" else "non_detect" if text == "<dl" else "note_required"
    return {"value": None, "status": status, "raw": value}

def extract(source):
    workbook = openpyxl.load_workbook(source, read_only=False, data_only=True)
    sheet = workbook['PFAS Reference']
    records = []
    for row in range(12, 23):
        for col in range(6, 78, 3):
            species = str(sheet.cell(10, col).value or '').split(' (')[0].replace('Gen X', 'GenX')
            if not species: continue
            values = [sheet.cell(row, c).value for c in range(col, col + 3)]
            if all(v is None for v in values): continue
            citation = f"PFAS Reference!{sheet.cell(row, col).coordinate}:{sheet.cell(row, col + 2).coordinate}"
            records.append({"id": f"pfas-r{row}-c{col}", "species": species,
                "site": sheet.cell(row, 1).value, "series": numeric(sheet.cell(row, 2).value),
                "totalEbctMinutes": numeric(sheet.cell(row, 3).value), "influentNgL": numeric(values[0]),
                "influentRaw": values[0], "tocMgL": numeric(sheet.cell(row, 5).value),
                "tocRaw": sheet.cell(row, 5).value,
                "bv1pct": observation(values[1]), "bv10pct": observation(values[2]),
                "bvDefinition": "total series media volume", "notes": sheet.cell(row, 78).value,
                "contextCitation": f"PFAS Reference!A{row}:E{row}", "notesCitation": f"PFAS Reference!BZ{row}",
                "citation": citation, "workbook": source.name,
                "waterType": "landfill leachate" if row in [13, 14] else "not reported",
                "media": "bituminous coal" if row == 19 else "coconut shell" if row == 20 else "Calgon F400" if row in [21,22] else "not reported"})
    guidance = [{"id": f"guidance-{i}", "citation": f"{name}!{cell}", "text": str(workbook[name][cell].value), "workbook": source.name}
        for i, (name, cell) in enumerate([('INPUT','I36'), ('INPUT','I46'), ('INPUT','I49'), ('Critical Design Assumptions','E69'), ('PFAS Reference','A3'), ('PFAS Reference','A5'), ('PFAS Reference','A6'), ('PFAS Reference','A7')])]
    return {"schemaVersion": "gac-knowledge/1", "source": {"file": source.name,
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "release": "March 2023",
        "extraction": "cached workbook values; no macros or recalculation", "path": f"docs/{source.name}"},
        "records": records, "guidance": guidance}

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, default=Path('docs/granular-activated-carbon-gac-.xlsm.xlsm'))
    parser.add_argument('--output', type=Path, default=Path('public/knowledge/gac_knowledge.json'))
    args = parser.parse_args()
    data = extract(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"Extracted {len(data['records'])} observations and {len(data['guidance'])} guidance records.")
