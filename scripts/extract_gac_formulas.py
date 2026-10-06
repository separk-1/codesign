"""Read formulas, named tables and cached validation values without changing Excel."""
import hashlib
import json
from pathlib import Path
import openpyxl

source = Path('docs/granular-activated-carbon-gac-.xlsm.xlsm')
formulas = openpyxl.load_workbook(source, data_only=False)
values = openpyxl.load_workbook(source, data_only=True)
cells = {}
for sheet, rows in {
    'Contactor Constraints': list(range(6, 15)) + list(range(18, 54)) + list(range(80, 84)),
    'Backwash and Regeneration': [13, 14, 31, 32, 33, 34, 35, 36, 38, 43],
    'Pumps Pipe Structure': [23, 24, 26, 27, 29, 30],
    'AutoSize': [15, 18, 23, 24, 58, 59, 60, 61, 62, 82, 83, 84, 85, 88, 89, 90, 91, 92, 93, 94, 95, 98, 99, 100, 101, 102, 103],
}.items():
    for row in rows:
        for col in ([3, 5] if sheet == 'AutoSize' else [3]):
            cell = formulas[sheet].cell(row, col)
            if cell.data_type == 'f':
                cells[f'{sheet}!{cell.coordinate}'] = {'formula': cell.value, 'cached': values[sheet][cell.coordinate].value}
tables = {}
for name in ['pipe_size_table_cl', 'vessel_size_table_cl']:
    sheet, region = next(formulas.defined_names[name].destinations)
    tables[name] = {'citation': f'{sheet}!{region}', 'rows': [[c.value for c in r] for r in values[sheet][region]]}
assumptions = {}
for row in values['Critical Design Assumptions'].iter_rows():
    if isinstance(row[1].value, str) and not row[1].value.startswith('='):
        assumptions[row[1].value] = {'value': row[2].value, 'citation': f'Critical Design Assumptions!{row[2].coordinate}'}
data = {'version': 'epa-wbs-gac-pressure/1', 'source': {'file': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest()},
    'formulas': cells, 'tables': tables, 'assumptions': assumptions,
    'baseline': {'contaminant': values['OUTPUT']['C3'].value, 'flowGpm': values['Contactor Constraints']['C8'].value,
                 'averageFlowGpm': values['Contactor Constraints']['C9'].value, 'totalEbctMinutes': values['INPUT']['G46'].value,
                 'series': values['INPUT']['G49'].value, 'diameterFt': values['INPUT']['G60'].value,
                 'heightFt': values['INPUT']['G59'].value, 'bedDepthFt': values['INPUT']['G57'].value,
                 'densityLbFt3': values['Critical Design Assumptions']['C70'].value,
                 'carbonLifeMode': 'bv', 'carbonLifeValue': values['INPUT']['G37'].value,
                 'bvDefinition': 'per_vessel', 'redundantVessels': values['INPUT']['G90'].value}}
out = Path('src/data/gacWorkbook.json'); out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'Extracted {len(cells)} formula cells and {len(tables)} lookup tables.')
